"""會議紀錄匯出成 Word（.docx）與 PDF。

Markdown 與 CSV 對工程師很好用，但要寄給老師、客戶或主管的會議紀錄，
對方期待的是打得開、印得出來的檔案。

兩個套件都只在真的要匯出時才 import：它們（尤其 reportlab 與 lxml）
加起來要幾秒才載入得完，沒有理由讓每次服務啟動都付這個成本。
"""
from __future__ import annotations

import io

_PRIORITY_ZH = {"high": "高", "medium": "中", "low": "低"}
_STATUS_ZH = {"todo": "待辦", "doing": "進行中", "done": "完成"}

PDF_FONT = "CJK"

# PDF 的中文字型一定要「內嵌」。reportlab 內建的 Adobe-CNS1 CID 字型
# （MSung-Light）不必附字型檔，看起來很方便，但它不內嵌字型資料——檢視器
# 只好拿手邊的字型去對 CID，結果是版面完全正確、每一個字都是別的字
# （「專題進度會議」變成「菡斌砦東滘琡」）。這種錯連「有沒有中文」的檢查
# 都過得了，只有真的把 PDF 畫出來才看得到。
#
# 所以改成找一份系統上的 CJK 字型內嵌進去。Linux 那幾個路徑由 Dockerfile 的
# fonts-arphic-uming 提供（繁體、而且是 TrueType 外框），Windows 的是給開發機用的。
#
# 為什麼不是 fonts-noto-cjk：那包是 OpenType/CFF 外框，reportlab 的 TTFont
# 只讀得懂 TrueType 的 glyf，載入會直接失敗。所以下面不是「挑一個存在的檔案」
# 而是「一個一個真的試著註冊，成功才算數」——光看路徑在不在是不夠的。
CJK_FONT_PATHS = (
    "/usr/share/fonts/truetype/arphic/uming.ttc",
    "/usr/share/fonts/truetype/arphic/ukai.ttc",
    "/usr/share/fonts/truetype/wqy/wqy-zenhei.ttc",
    "/usr/share/fonts/truetype/wqy/wqy-microhei.ttc",
    "/usr/share/fonts/truetype/droid/DroidSansFallbackFull.ttf",
    "/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc",
    r"C:\Windows\Fonts\msjh.ttc",
    r"C:\Windows\Fonts\mingliu.ttc",
    r"C:\Windows\Fonts\kaiu.ttf",
    "/System/Library/Fonts/PingFang.ttc",
)
# .ttc 裡有好幾套字（日／韓／簡／繁）。字都畫得出來，但「骨、直、戶」這些
# 字形依地區而異，所以優先挑繁體那一套
_TC_HINTS = ("tc", "tw", "jheng", "ming", "hk", "hei")


class PdfFontMissing(RuntimeError):
    """這台機器上找不到任何 reportlab 載得進來的中文字型，產不出看得懂的 PDF。"""

TASK_HEADERS = ["任務", "負責人", "期限", "優先級", "狀態"]


def _task_row(t: dict) -> list[str]:
    return [
        t.get("task", ""),
        t.get("owner") or "未指派",
        t.get("due_date") or "未定",
        _PRIORITY_ZH.get(t.get("priority"), ""),
        _STATUS_ZH.get(t.get("status"), ""),
    ]


def _blocks(record: dict, tasks: list[dict]) -> list[tuple[str, object]]:
    """兩種格式共用的內容骨架：[(型別, 內容), ...]。

    排版差很多，但「一份會議紀錄裡有什麼、順序是什麼」只該寫一次——
    不然加一個區塊時一定會漏掉其中一種格式。
    """
    m = record.get("meeting") or {}
    out: list[tuple[str, object]] = [
        ("title", m.get("title") or "（未命名會議）"),
        ("meta", f"日期：{m.get('date') or '—'}　|　出席：{'、'.join(m.get('attendees') or []) or '（未記錄）'}"),
        ("h", "摘要"),
        ("p", m.get("summary") or "（未產生摘要）"),
        ("h", "決議"),
    ]
    decisions = record.get("decisions") or []
    out += [("li", d["description"] + (f"（{d['context']}）" if d.get("context") else ""))
            for d in decisions] or [("p", "（本次會議無正式決議）")]

    highlights = record.get("highlights") or []
    if highlights:
        out.append(("h", "會議重點"))
        out += [("li", h["text"] + (f"（{h['time']}）" if h.get("time") else "")) for h in highlights]

    # 自訂範本的區塊排在代辦前面：它們是那種會議真正要看的東西
    for sec in record.get("sections") or []:
        out.append(("h", sec.get("label") or ""))
        out += [("li", item) for item in sec.get("items") or []] or [("p", "（本次未提及）")]

    out.append(("h", "代辦事項"))
    if tasks:
        out.append(("table", [TASK_HEADERS, *(_task_row(t) for t in tasks)]))
    else:
        out.append(("p", "（無代辦事項）"))

    pending = record.get("pending_items") or []
    if pending:
        out.append(("h", "待決事項"))
        out += [("li", p["topic"] + (f"（{p['reason']}）" if p.get("reason") else "")) for p in pending]
    return out


def meeting_to_docx(record: dict, tasks: list[dict]) -> bytes:
    from docx import Document
    from docx.shared import Pt

    doc = Document()
    for kind, content in _blocks(record, tasks):
        if kind == "title":
            doc.add_heading(str(content), level=0)
        elif kind == "meta":
            run = doc.add_paragraph().add_run(str(content))
            run.italic = True
            run.font.size = Pt(9)
        elif kind == "h":
            doc.add_heading(str(content), level=1)
        elif kind == "p":
            doc.add_paragraph(str(content))
        elif kind == "li":
            doc.add_paragraph(str(content), style="List Bullet")
        elif kind == "table":
            rows = content
            table = doc.add_table(rows=len(rows), cols=len(rows[0]))
            table.style = "Table Grid"
            for r, row in enumerate(rows):
                for c, cell in enumerate(row):
                    table.cell(r, c).text = str(cell)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _faces(path: str) -> list[int]:
    """這個字型檔裡有哪幾套字，繁體那一套排前面。

    .ttc 常常把日／韓／簡／繁包在一起。字都畫得出來，但「骨、直、戶」這些
    字形依地區而異，所以優先挑繁體。
    """
    if not path.lower().endswith(".ttc"):
        return [0]

    from reportlab.pdfbase.ttfonts import TTFontFile

    preferred, rest = [], []
    for index in range(16):
        try:
            name = TTFontFile(path, subfontIndex=index, validate=0).name.decode("latin-1").lower()
        except Exception:  # noqa: BLE001 - 索引超出範圍就是這套字型只有這麼多
            break
        (preferred if any(h in name for h in _TC_HINTS) else rest).append(index)
    return preferred + rest or [0]


def find_cjk_font(paths=CJK_FONT_PATHS) -> tuple[str, int]:
    """回傳（字型檔路徑, .ttc 裡的第幾套）；一個都載不進來就丟 PdfFontMissing。

    刻意「真的試著載一次」而不是只看檔案在不在：fonts-noto-cjk 那種
    OpenType/CFF 外框路徑存在、檔案也正常，但 reportlab 的 TTFont 只讀得懂
    TrueType 的 glyf，載入會直接失敗。光檢查路徑的話，會在產 PDF 的當下才炸。
    """
    import os

    from reportlab.pdfbase.ttfonts import TTFont

    for path in paths:
        if not os.path.exists(path):
            continue
        for index in _faces(path):
            try:
                TTFont("probe", path, subfontIndex=index)
            except Exception:  # noqa: BLE001 - 載不進來就換下一套／下一個檔案
                continue
            return path, index
    raise PdfFontMissing(
        "伺服器上找不到 reportlab 載得進來的中文字型，無法產生 PDF"
        "（Docker 映像需要 fonts-arphic-uming 這類 TrueType 外框的中文字型）"
    )


def register_pdf_font() -> str:
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont

    if PDF_FONT not in pdfmetrics.getRegisteredFontNames():
        path, index = find_cjk_font()
        pdfmetrics.registerFont(TTFont(PDF_FONT, path, subfontIndex=index))
    return PDF_FONT


def meeting_to_pdf(record: dict, tasks: list[dict]) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import ListFlowable, ListItem, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    register_pdf_font()
    title_style = ParagraphStyle("zh-title", fontName=PDF_FONT, fontSize=18, leading=26, spaceAfter=4)
    meta_style = ParagraphStyle("zh-meta", fontName=PDF_FONT, fontSize=9, leading=14,
                                textColor=colors.grey, spaceAfter=10)
    h_style = ParagraphStyle("zh-h", fontName=PDF_FONT, fontSize=13, leading=20,
                             spaceBefore=12, spaceAfter=4)
    body = ParagraphStyle("zh-body", fontName=PDF_FONT, fontSize=10.5, leading=17)

    def escape(text: str) -> str:
        # Paragraph 吃的是小型標記語言，內容裡的 < & 會被當成標籤
        return str(text).replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")

    flow, bullets = [], []

    def flush_bullets() -> None:
        if bullets:
            flow.append(ListFlowable(
                [ListItem(Paragraph(escape(b), body), leftIndent=14) for b in bullets],
                bulletType="bullet", bulletFontName=PDF_FONT, bulletFontSize=8, leftIndent=14,
            ))
            bullets.clear()

    for kind, content in _blocks(record, tasks):
        if kind != "li":
            flush_bullets()
        if kind == "title":
            flow.append(Paragraph(escape(content), title_style))
        elif kind == "meta":
            flow.append(Paragraph(escape(content), meta_style))
        elif kind == "h":
            flow.append(Paragraph(escape(content), h_style))
        elif kind == "p":
            flow.append(Paragraph(escape(content), body))
        elif kind == "li":
            bullets.append(content)
        elif kind == "table":
            rows = [[Paragraph(escape(c), body) for c in row] for row in content]
            table = Table(rows, colWidths=[62 * mm, 24 * mm, 26 * mm, 18 * mm, 18 * mm], repeatRows=1)
            table.setStyle(TableStyle([
                ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#c9c9c9")),
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#f0f0f0")),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("TOPPADDING", (0, 0), (-1, -1), 4),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
            ]))
            flow += [Spacer(1, 4), table]
    flush_bullets()

    buf = io.BytesIO()
    SimpleDocTemplate(
        buf, pagesize=A4, title=(record.get("meeting") or {}).get("title") or "會議紀錄",
        leftMargin=18 * mm, rightMargin=18 * mm, topMargin=18 * mm, bottomMargin=18 * mm,
    ).build(flow)
    return buf.getvalue()
